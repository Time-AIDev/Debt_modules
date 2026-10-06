#!/usr/bin/env python3
"""Render an evidence-scoped DOCX and public aggregate JSON, never raw inputs.

Private source files are read only. Unknown fields/text are discarded by an
explicit whitelist; neither a progress claim nor a static inventory certifies
deployment, money, capacity or usability acceptance.
"""
import argparse
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

REVISION = re.compile(r"[0-9a-f]{40}\Z")
SHA = re.compile(r"[0-9a-f]{64}\Z")
IMAGE = re.compile(r"sha256:[0-9a-f]{64}\Z")
THRESHOLDS = {"business_read_ms": "p(95)<2000", "business_write_ms": "p(95)<2000",
              "business_errors": "rate<0.01", "checks": "rate>0.99", "idempotency_replay_errors": "rate==0"}
SUPPLEMENTAL = {"paged": {"capacity_page_ms": "p(95)<2000", "capacity_page_errors": "rate<0.01"},
                "debt-workflow": {"complete_action_ms": "p(95)<2000", "debt_mutations_committed": "count>0",
                                  "role_denial_errors": "rate==0"}}
INPUTS = ("S3 configuration", "age recipient/key custodian", "recovery host", "independent load generator",
          "HTTP JSON webhook configuration", "Git write credentials")
ROLES = {"OWNER", "MANAGER", "SUPERVISOR", "OFFICER", "OPS"}
UX_DIMENSIONS = {"completion", "navigation", "financial_clarity", "recovery", "accessibility"}


def read_json(path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return None


def identifier(value, pattern):
    return value if isinstance(value, str) and pattern.fullmatch(value) else None


def number(value, *, integer=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        return None
    return int(value) if integer and int(value) == value else None if integer else float(value)


def amount(value):
    if not isinstance(value, str):
        return None
    try:
        parsed = Decimal(value)
        return str(parsed.quantize(Decimal("0.01"))) if parsed.is_finite() else None
    except InvalidOperation:
        return None


def date_only(value):
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        try:
            datetime.strptime(value, "%Y-%m-%d")
            return value
        except ValueError:
            pass
    return None


def financial(source):
    out = {}
    for kind in ("invoice", "receipt"):
        item = (source or {}).get(kind) or {}
        row = {key: number(item.get(key), integer=True) for key in (
            "documents", "components", "multi_component_documents", "source_rows", "rejected_rows",
            "quarantined_documents", "signed_adjustment_documents")}
        row.update(source_sha256=identifier(item.get("sha256"), SHA), freshness=date_only(item.get("freshness")),
                   parsed_component_total=amount(item.get("component_total")))
        out[kind] = row
    return {"source_validation": out, "applied_or_allocated_cash_certified": False,
            "scope": "Parser/component validation only; totals do not establish collected cash or every case mapping"}


def staged_financial(source, revision, *, expected_api_image=None):
    """Publish aggregate staged observations without copying customer-level proof."""
    source = source or {}
    observed_revision = identifier(source.get("revision"), REVISION)
    runtime = source.get("runtime_kind") if source.get("runtime_kind") in ("ASGI", "HTTP_IMAGE") else None
    api_image = identifier(source.get("api_image"), IMAGE)
    hashes = {kind: identifier((source.get("source_sha256") or {}).get(kind), SHA)
              for kind in ("master", "invoice", "receipt")}
    check_names = ("source_rows_accounted", "net_cash_conserved", "snapshot_cash_not_double_counted",
                   "duplicate_import_once", "operator_work_preserved", "all_mapped_cases_readable",
                   "cancellation_evidence_verified", "manual_npl_retained", "installation_unknown_held")
    checks = {name: (source.get("checks") or {}).get(name) if isinstance(
        (source.get("checks") or {}).get(name), bool) else None for name in check_names}
    count_names = ("customer_count", "contracts_applied", "status_observations_applied", "terminations_applied",
                   "installation_confirmations_applied", "baseline_person_entries", "retained_person_entries",
                   "unworked_pipeline_seeds_pruned", "lost_operator_entries", "served_financial_cases",
                   "served_case_errors", "accepted_invoice_rows", "accepted_receipt_rows", "quarantined_receipt_rows",
                   "quarantined_receipt_documents", "eligible_paid_receipts_conserved", "paid_receipts_held_for_review",
                   "cancellation_documents_reversed")
    counts = {name: number((source.get("counts") or {}).get(name), integer=True) for name in count_names}
    totals = {name: amount((source.get("totals") or {}).get(name)) for name in (
        "allocated", "customer_credit", "eligible_receipt_net", "review_receipt_net", "cancel_reversed_net")}
    timings = {name: number((source.get("timings_ms") or {}).get(name)) for name in (
        "odoo_pull", "supplement_preview", "supplement_apply", "status_preview", "status_apply",
        "invoice_enqueue", "invoice_apply", "receipt_enqueue", "receipt_apply", "duplicate_replay")}
    completed = (source.get("format") == "volt-staging-financial-execution-v1"
                 and source.get("execution_status") == "completed")
    verified = (completed and observed_revision == revision and runtime == "HTTP_IMAGE" and api_image is not None
                and api_image == identifier(expected_api_image, IMAGE)
                and all(hashes.values()) and all(value is True for value in checks.values())
                and source.get("production_mutated") is False)
    return {"observed_revision": observed_revision, "matching_report_revision": observed_revision == revision,
            "execution_completed": completed, "runtime_kind": runtime, "api_image": api_image,
            "source_sha256": hashes, "checks": checks, "counts": counts, "totals": totals, "timings_ms": timings,
            "staged_workflow_verified_for_report_revision": bool(verified), "production_application_certified": False,
            "scope": "Aggregate staged observations; source-only or older runtime is not final-image acceptance"}


def qa_aggregate(matrix, scenarios, execution, revision):
    flows = (matrix or {}).get("flows") or []
    valid_flows = [flow for flow in flows if isinstance(flow, dict)]
    statuses = Counter()
    allowed = {"passed", "failed", "skipped", "pending", "not_executed", "inventory_only", "not_observed"}
    for flow in valid_flows:
        state = flow.get("status")
        statuses[state if state in allowed else "inventory_only"] += 1
    scenarios = scenarios if isinstance(scenarios, list) else []
    execution = execution if isinstance(execution, dict) else {}
    observed_dimensions = execution.get("dimensions")
    observed_dimensions = observed_dimensions if isinstance(observed_dimensions, dict) else {}
    dimensions = {name: number(observed_dimensions.get(name)) for name in sorted(UX_DIMENSIONS)}
    observed_roles = execution.get("roles_observed")
    roles = sorted(role for role in ROLES if isinstance(observed_roles, list) and role in observed_roles)
    completion = {key: number(execution.get(key), integer=True) for key in (
        "feature_count", "checkpoints_total", "checkpoints_passed", "open_defects")}
    score = None
    executed = bool(execution and execution.get("execution_status") == "completed"
                    and execution.get("revision") == revision)
    if executed:
        candidate = number(execution.get("score"))
        if candidate is not None and candidate <= 10 and number(execution.get("unexecuted"), integer=True) == 0:
            score = candidate
            if candidate == 10 and not (
                number(execution.get("failed"), integer=True) == 0
                and number(execution.get("skipped"), integer=True) == 0
                and number(execution.get("open_defects"), integer=True) == 0
                and (number(execution.get("passed"), integer=True) or 0) > 0
                and (completion["feature_count"] or 0) >= 19
                and (completion["checkpoints_total"] or 0) >= 33
                and completion["checkpoints_passed"] == completion["checkpoints_total"]
                and set(observed_dimensions) == UX_DIMENSIONS
                and all(value == 2 for value in dimensions.values())
                and set(roles) >= {"OFFICER", "MANAGER", "OPS"}
            ):
                score = None
    counts = {key: number((execution or {}).get(key), integer=True) if executed else None
              for key in ("passed", "failed", "skipped", "unexecuted")}
    return {"inventory_revision": identifier((matrix or {}).get("revision"), REVISION),
            "inventory_flows": len(valid_flows), "inventory_routes": len((matrix or {}).get("routes") or []),
            "inventory_scenarios": len(scenarios), "inventory_mutation_scenarios": sum(
                row.get("mutation") is True for row in scenarios if isinstance(row, dict)),
            "browser_observed_flows": sum(flow.get("browser_observed") is True for flow in valid_flows),
            "mutation_observed_flows": sum(flow.get("mutation_observed") is True for flow in valid_flows),
            "inventory_status_counts": dict(statuses), "execution_verified_for_report_revision": executed,
            "execution_counts": counts, "usability_estimate_score": score,
            "execution_completion": completion if executed else None,
            "execution_dimensions": dimensions if executed else None,
            "roles_observed": roles if executed else [],
            "score_scope": "Revision-bound estimate, not human approval" if score is not None else "Not scored: execution evidence incomplete",
            "inventory_is_execution": False}


def test_aggregate(path, revision):
    data = read_json(path) or {}
    components = {}
    for name in ("backend", "frontend"):
        value = data.get(name)
        if not isinstance(value, dict):
            continue
        components[name] = {key: number(value.get(key), integer=True) for key in ("tests", "passed", "failed", "skipped")}
    def manifest(value):
        value = value if isinstance(value, dict) else {}
        return {"revision": identifier(value.get("revision"), REVISION),
                "tracked_sha256": identifier(value.get("tracked_sha256"), SHA),
                "tracked_clean": value.get("tracked_clean") if isinstance(value.get("tracked_clean"), bool) else None,
                "untracked_code_files": number(value.get("untracked_code_files"), integer=True)}
    before, after = (manifest(data.get(key)) for key in ("source_before", "source_after"))
    unchanged = data.get("source_unchanged") if isinstance(data.get("source_unchanged"), bool) else None
    verified = (unchanged is True and before == after and before["revision"] is not None
                and before["tracked_sha256"] is not None and before["tracked_clean"] is True
                and before["untracked_code_files"] == 0)
    matching = before["revision"] == revision and after["revision"] == revision
    binding = ("MATCHING_REPORT_REVISION" if verified and matching else "OLDER_REVISION" if verified else
               "MUTABLE_SOURCE" if unchanged is False or before["tracked_clean"] is False or after["tracked_clean"] is False
               else "UNBOUND")
    return {"components": components, "production_used": data.get("production_used") is True,
            "source_before": before, "source_after": after, "source_unchanged": unchanged,
            "matching_report_revision": matching, "source_bound_for_report_revision": bool(verified and matching),
            "binding": binding, "final_regression_certified": False,
            "scope": "Source-bound component execution; whole-suite acceptance requires final regression gates"
                if verified and matching else "Older, mutable or unbound execution; not final regression evidence",
            "counts_are_not_summed_across_overlapping_suites": True}


def k6_runs(directory, runtime, revision):
    runs = []
    runtime_revision = identifier((runtime or {}).get("revision"), REVISION)
    api_image = identifier((runtime or {}).get("api_image"), IMAGE)
    for summary in sorted(directory.glob("k6-*-summary.json")):
        match = re.fullmatch(r"k6-([0-9a-f]{7,40})-(historical|paged|debt-workflow)-(\d+)vu"
                            r"(?:-(mixed|import|growth|soak))?-summary\.json", summary.name)
        if not match:
            continue
        meta_path = summary.with_name(summary.name.replace("-summary.json", "-meta.json"))
        data, meta = read_json(summary) or {}, read_json(meta_path) or {}
        full = identifier(meta.get("revision"), REVISION)
        if (full is None and "revision" not in meta and "api_image" not in meta
                and runtime_revision and runtime_revision.startswith(match[1])):
            full = runtime_revision
        measured_image = identifier(meta.get("api_image"), IMAGE)
        if measured_image is None and full == runtime_revision:
            measured_image = api_image
        bound = full is not None and full.startswith(match[1]) and measured_image is not None
        def measured(rules):
            result = {}
            for name, rule in rules.items():
                native = (data.get("metrics") or {}).get(name) or {}
                values = native.get("values") or {}
                observed = number(values.get("p(95)" if name.endswith("_ms") else "count" if rule.startswith("count") else "rate"))
                native_gate = (native.get("thresholds") or {}).get(rule, {}).get("ok")
                result[name] = {"value": observed, "threshold": rule,
                               "native_threshold_passed": native_gate if isinstance(native_gate, bool) else None}
            return result
        metrics = measured(THRESHOLDS)
        supplemental = measured(SUPPLEMENTAL.get(match[2], {}))
        passed = all(value["native_threshold_passed"] is True and value["value"] is not None for value in metrics.values())
        supplemental_passed = all(value["native_threshold_passed"] is True and value["value"] is not None
                                  for value in supplemental.values()) if supplemental else None
        preview = meta.get("concurrent_preview")
        safe_preview = None
        if isinstance(preview, dict):
            safe_preview = {key: preview.get(key) is True for key in ("started", "passed", "never_applied", "isolated_staging_only")}
            safe_preview.update(profile=preview.get("profile") if preview.get("profile") in ("debt_status", "contract_supplement") else None,
                status=number(preview.get("status"), integer=True), elapsed_ms=number(preview.get("elapsed_ms")),
                source_bytes=number(preview.get("source_bytes"), integer=True), source_sha256=identifier(preview.get("source_sha256"), SHA))
        health_guard = meta.get("production_health_guard")
        safe_guard = None
        if isinstance(health_guard, dict):
            reasons = {"production_health", "host_low_memory", "target_oom", "target_restart", "target_not_running", "guard_unavailable", "guard_monitor_incomplete"}
            safe_guard = {key: health_guard.get(key) is True for key in ("enabled", "baseline_passed", "aborted", "isolated_api_stopped", "monitor_completed", "shutdown_complete")}
            safe_guard.update(reason=health_guard.get("reason") if health_guard.get("reason") in reasons else None,
                samples=len(health_guard.get("samples") or []),
                minimum_available_memory_bytes=number(health_guard.get("minimum_available_memory_bytes"), integer=True))
        dataset = meta.get("dataset") or {}
        roles = {role: number(count, integer=True) for role, count in (meta.get("actor_roles") or {}).items() if role in ROLES}
        runs.append({"profile": match[2], "scenario": match[4], "vus": int(match[3]), "duration_seconds": number(meta.get("duration_seconds"), integer=True),
            "revision": full, "api_image": measured_image if bound else None,
            "binding": "MATCHING_REPORT_REVISION" if bound and full == revision else "OLDER_REVISION" if bound else "UNBOUND",
            "classification": "COLOCATED_PROVISIONAL" if meta.get("generator_co_located") is True else "GENERATOR_LOCATION_UNVERIFIED",
            "native_exit_code": number(meta.get("process_return_code"), integer=True), "original_thresholds_passed": passed,
            "orchestration_exit_code": number(meta.get("exit_code"), integer=True),
            "capacity_acceptance_certified": False, "metrics": metrics, "supplemental_metrics": supplemental,
            "supplemental_thresholds_passed": supplemental_passed, "concurrent_preview": safe_preview,
            "production_health_guard": safe_guard,
            "diagnostic_run_passed": passed and supplemental_passed is not False and meta.get("process_return_code") == 0
                and meta.get("exit_code", 0) == 0 and (safe_preview is None or all(safe_preview[key] for key in ("started", "passed", "never_applied", "isolated_staging_only")))
                and (safe_guard is None or safe_guard["baseline_passed"] and not safe_guard["aborted"] and safe_guard["monitor_completed"]),
            "dataset": {key: number(dataset.get(key), integer=True) for key in ("customers", "contracts", "instalments_per_contract")},
            "actor_roles": roles, "workload_sha256": identifier(meta.get("workload_sha256"), SHA),
            "historical_workload_sha256": identifier(meta.get("historical_workload_sha256"), SHA),
            "summary_sha256": hashlib.sha256(summary.read_bytes()).hexdigest(),
            "meta_sha256": hashlib.sha256(meta_path.read_bytes()).hexdigest() if meta_path.exists() else None})
    return runs


def build_aggregate(evidence, revision, *, qa_execution=None, browser=None, test_summaries=()):
    progress = read_json(evidence / "progress.json") or {}
    runtime = read_json(evidence / "diagnostic/owned-resources.json") or {}
    google = "USER_ATTESTED_SUCCESS"
    if (browser and browser.get("revision") == revision and browser.get("browser_login_performed") is True
            and browser.get("authenticated_google_session") is True and browser.get("authenticated_session_kind") == "GOOGLE"
            and browser.get("status") == 200 and browser.get("google_origin_denied") is False
            and browser.get("target") == "https://collection.volt.co.th" and browser.get("failed_requests") == []
            and browser.get("page_error_count") == 0):
        google = "FINAL_CANDIDATE_BROWSER_VERIFIED"
    pending = [item for item in INPUTS if item in (progress.get("external_inputs_pending") or [])]
    return {"format": "volt-implementation-aggregate-v1", "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "user_timezone": "Asia/Bangkok", "report_revision": revision,
        "manifest_revision": identifier(progress.get("committed_revision") or progress.get("integration_head"), REVISION),
        "scope": "Implementation evidence snapshot; completion/deployment not certified",
        "completion_certified": False, "deployment_certified": False, "capacity_acceptance_certified": False,
        "google_login_status": google, "go_live_enabled_reported": progress.get("go_live_enabled") is True,
        "offsite_restore_verified_reported": progress.get("offsite_restore_verified") is True,
        "pending_external_inputs": pending,
        "original_k6_thresholds": THRESHOLDS,
        "k6": k6_runs(evidence / "diagnostic/results", runtime, revision),
        "financial": {**financial(read_json(evidence / "component-real-file-validation.json")),
                      "staging_execution": staged_financial(read_json(evidence / "financial-execution.json"), revision,
                          expected_api_image=progress.get("candidate_api_image") if progress.get("candidate_revision") == revision else None)},
        "ux": qa_aggregate(read_json(evidence / "qa/feature-matrix.json"),
                           read_json(evidence / "qa/scenarios.template.json"), qa_execution, revision),
        "tests": [test_aggregate(path, revision) for path in test_summaries]}


def table(doc, columns, rows):
    grid = doc.add_table(rows=1, cols=len(columns))
    grid.style = "Light Shading Accent 1"
    for cell, heading in zip(grid.rows[0].cells, columns):
        cell.text = heading
    for row in rows:
        for cell, value in zip(grid.add_row().cells, row):
            cell.text = "—" if value is None else str(value)


def render_docx(data, output):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.section import WD_ORIENT
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.left_margin = section.right_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Tahoma"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("VOLT — รายงานความคืบหน้าการปรับปรุงระบบ", 0)
    stamp = datetime.fromisoformat(data["generated_at_utc"]).astimezone(__import__("zoneinfo").ZoneInfo("Asia/Bangkok"))
    doc.add_paragraph("ข้อมูล ณ %s (Asia/Bangkok) · revision %s" % (stamp.strftime("%Y-%m-%d %H:%M"), data["report_revision"]))
    doc.add_heading("Readiness และขอบเขตหลักฐาน", 1)
    doc.add_paragraph("เอกสารนี้รายงานหลักฐานการพัฒนาและทดสอบ ไม่รับรองว่างานทั้งหมดเสร็จ เปิด go-live หรือ deploy production แล้ว")
    doc.add_paragraph("Google login: " + data["google_login_status"] + ". คำยืนยันจากผู้ใช้แยกจากการตรวจ browser บน revision สุดท้าย")
    doc.add_paragraph("ผล load บน VPS เดียวกับตัวสร้างโหลดเป็น COLOCATED_PROVISIONAL จึงใช้วิเคราะห์คอขวด และยังไม่ใช่หลักฐานยอมรับ capacity จากเครื่องอิสระ")
    doc.add_heading("k6 — เกณฑ์เดิมและผลจริง", 1)
    doc.add_paragraph("อ่าน/เขียน p95 <2,000 ms; business errors <1%; checks >99%; idempotency replay errors =0. ไม่มีการผ่อนเกณฑ์")
    rows = []
    for run in data["k6"]:
        m = run["metrics"]
        def fmt(name, rate=False):
            value = m[name]["value"]
            return None if value is None else ("%.3f%%" % (value * 100) if rate else "%.2f" % value)
        label = run["profile"] + (" / " + run["scenario"] if run["scenario"] else "")
        rows.append((label, run["vus"], fmt("business_read_ms"), fmt("business_write_ms"),
                     fmt("business_errors", True), fmt("checks", True), fmt("idempotency_replay_errors", True),
                     "PASS" if run["original_thresholds_passed"] else "FAIL / incomplete", run["binding"]))
    table(doc, ("Profile", "VUs", "Read p95 ms", "Write p95 ms", "Error rate", "Checks rate", "Replay rate", "Original gates", "Revision binding"), rows)
    for run in data["k6"]:
        doc.add_paragraph("%s / %s VUs: revision %s · API %s · %s · exit %s" %
                          (run["profile"], run["vus"], run["revision"], run["api_image"], run["classification"], run["native_exit_code"]))
        doc.add_paragraph("Combined diagnostic gates: " + ("PASS" if run["diagnostic_run_passed"] else "FAIL / incomplete") +
                          "; orchestration exit=" + str(run["orchestration_exit_code"]))
        if run["supplemental_metrics"]:
            doc.add_paragraph("Supplemental native gates: " + ("PASS" if run["supplemental_thresholds_passed"] else "FAIL / incomplete") +
                "; " + "; ".join("%s=%s (%s)" % (key, value["value"], value["threshold"]) for key, value in run["supplemental_metrics"].items()))
        if run["concurrent_preview"]:
            preview = run["concurrent_preview"]
            doc.add_paragraph("Concurrent preview: %s · status %s · %s ms · %s · never applied=%s" %
                (preview["profile"], preview["status"], preview["elapsed_ms"], "PASS" if preview["passed"] else "FAIL / incomplete", preview["never_applied"]))
        if run["production_health_guard"]:
            guard = run["production_health_guard"]
            doc.add_paragraph("Production/host guard: baseline=%s · aborted=%s · reason=%s · observations=%s" %
                (guard["baseline_passed"], guard["aborted"], guard["reason"], guard["samples"]))
    doc.add_heading("ข้อมูลเงินจริงและการตรวจไฟล์", 1)
    doc.add_paragraph("ตัวเลขต่อไปนี้เป็นผล parser/component validation ของไฟล์ที่ได้รับ ไม่ใช่ยอดเงินเข้าที่จัดสรรแล้ว และไม่ได้ยืนยัน mapping ทุก customer case")
    table(doc, ("Source", "Documents", "Components", "Multi-payment docs", "Rejected rows", "Quarantined docs", "Parsed component total", "As of"),
          [(kind, row["documents"], row["components"], row["multi_component_documents"], row["rejected_rows"],
            row["quarantined_documents"], row["parsed_component_total"], row["freshness"])
           for kind, row in data["financial"]["source_validation"].items()])
    doc.add_paragraph("หลักฐานตรวจไฟล์ไม่รับรองว่าได้ apply ข้อมูลบน production; เคสขัดแย้งต้องผ่านการ review ตาม workflow เดิม")
    stage = data["financial"]["staging_execution"]
    doc.add_paragraph("Staging workflow: revision %s · runtime %s · final-image proof %s" %
                      (stage["observed_revision"], stage["runtime_kind"],
                       "PASS" if stage["staged_workflow_verified_for_report_revision"] else "incomplete / provisional"))
    table(doc, ("Staged financial check", "Observed result"), list(stage["checks"].items()))
    table(doc, ("Aggregate count", "Observed"), [(name, value) for name, value in stage["counts"].items() if value is not None])
    table(doc, ("Staged amount", "Observed net"), [(name, value) for name, value in stage["totals"].items() if value is not None])
    doc.add_paragraph("ยอด staging เป็นผลจำลองจากไฟล์จริงบนข้อมูลสำเนา ส่วนรายการที่ส่งตรวจยังไม่ใช้เป็นเงินรับชำระยืนยันบน production")
    doc.add_heading("Debt team UX และความครอบคลุม", 1)
    ux = data["ux"]
    doc.add_paragraph("Inventory: %s flows / %s routes / %s scenarios. Inventory และรายชื่อ tests แยกจากการรันทดสอบจริง" %
                      (ux["inventory_flows"], ux["inventory_routes"], ux["inventory_scenarios"]))
    doc.add_paragraph("Browser-observed flows: %s; mutation-observed flows: %s. คะแนน usability estimate: %s" %
                      (ux["browser_observed_flows"], ux["mutation_observed_flows"], ux["usability_estimate_score"] if ux["usability_estimate_score"] is not None else "ยังไม่ให้คะแนน"))
    table(doc, ("Passed", "Failed", "Skipped", "Unexecuted"), [[ux["execution_counts"][key] for key in ("passed", "failed", "skipped", "unexecuted")]])
    if ux["execution_completion"]:
        completion = ux["execution_completion"]
        doc.add_paragraph("Observed flows: %s; checkpoints: %s/%s; open defects: %s; roles: %s" %
            (completion["feature_count"], completion["checkpoints_passed"], completion["checkpoints_total"],
             completion["open_defects"], ", ".join(ux["roles_observed"])))
        table(doc, ("Usability dimension", "Observed score"), list(ux["execution_dimensions"].items()))
    doc.add_heading("การทดสอบ component ที่มีหลักฐาน execution", 1)
    rows = []
    for index, suite in enumerate(data["tests"], 1):
        for component, counts in suite["components"].items():
            rows.append((index, component, counts["tests"], counts["passed"], counts["failed"], counts["skipped"]))
    table(doc, ("Evidence set", "Component", "Tests", "Passed", "Failed", "Skipped"), rows)
    for index, suite in enumerate(data["tests"], 1):
        doc.add_paragraph("Evidence set %s: %s; source unchanged=%s; matching report revision=%s. %s" %
            (index, suite["binding"], suite["source_unchanged"], suite["matching_report_revision"], suite["scope"]))
        for side in ("source_before", "source_after"):
            manifest = suite[side]
            doc.add_paragraph("%s: revision %s; tracked SHA256 %s; tracked clean=%s; untracked code files=%s" %
                (side, manifest["revision"], manifest["tracked_sha256"], manifest["tracked_clean"], manifest["untracked_code_files"]))
    doc.add_paragraph("ไม่บวกจำนวนจากหลาย suite เป็น tests ไม่ซ้ำ เพราะชุดทดสอบอาจทับซ้อนกัน; ช่องว่างหมายถึงไม่มีหลักฐานตัวเลข")
    doc.add_heading("สิ่งที่ยังต้องยืนยันก่อน go-live", 1)
    for item in data["pending_external_inputs"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph("ต้องมี capacity บน generator อิสระ, workflow การเงินจริงครบวงจร, QA บน candidate สุดท้าย, offsite restore/alert/rollback และ Google browser evidence ที่ผูกกับ revision ตาม release gates")
    doc.save(output)
    with zipfile.ZipFile(output) as archive:
        if archive.testzip() is not None:
            raise ValueError("DOCX CRC validation failed")
        ET.fromstring(archive.read("word/document.xml"))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--revision", required=True)
    parser.add_argument("--test-summary", action="append", type=Path, default=[])
    parser.add_argument("--qa-execution", type=Path)
    parser.add_argument("--browser-evidence", type=Path)
    args = parser.parse_args(argv)
    if not identifier(args.revision, REVISION):
        parser.error("Use the full 40-character evidence/report revision")
    data = build_aggregate(args.evidence_dir, args.revision,
        qa_execution=read_json(args.qa_execution) if args.qa_execution else None,
        browser=read_json(args.browser_evidence) if args.browser_evidence else None,
        test_summaries=args.test_summary)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    aggregate, document = args.output_dir / "implementation-aggregate.json", args.output_dir / "implementation-readiness-and-k6.docx"
    aggregate.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    render_docx(data, document)
    print(json.dumps({"docx": str(document), "aggregate": str(aggregate), "docx_sha256": hashlib.sha256(document.read_bytes()).hexdigest(),
                      "capacity_acceptance_certified": False, "deployment_certified": False}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
