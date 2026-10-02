#!/usr/bin/env python3
"""Build the login tester guide (design 2026-09-13 §4) as two .docx files from one source.

  scripts/guides/.venv/bin/python scripts/guides/build_login_guide.py

  --repo-out   docs/guides/login-tester-guide.docx       (committed; password cells masked)
  --vault-out  <vault>/06 - Login Tester Guide (with passwords).docx
  --accounts   ~/.volt-test-accounts.json                (from make_test_accounts.py)
"""
import argparse
import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
import role_matrix  # noqa: E402

VERSION = "1.0"
VAULT = Path.home() / "Documents" / "Obsidian Vault" / "Debt modules (self build)"
FONT = "Sarabun"
MASK = "ดูฉบับในคลังเอกสาร (vault copy)"
ROLE_TH = {"OFFICER": "เจ้าหน้าที่", "SUPERVISOR": "หัวหน้างาน", "MANAGER": "ผู้จัดการ",
           "OWNER": "เจ้าของ", "OPS": "ทีมปฏิบัติการ (O&M)"}
ACCOUNTS = (
    ("tester", "ทุกบทบาท (สลับได้) · any role", "เลือกเอง · chooses"),
    ("officer.test", "OFFICER", "nantapol"),
    ("supervisor.test", "SUPERVISOR", "settasak"),
    ("manager.test", "MANAGER", "settasak"),
    ("owner.test", "OWNER", "owner"),
    ("ops.test", "OPS", "ops"),
)
AMBER = "F5A524"


# --- template ------------------------------------------------------------------------

def _font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, hex_fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def _page_number(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    _font(run, 9)


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.2))
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        style = doc.styles[name]
        style.font.name = FONT
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)
    doc.styles["Normal"].font.size = Pt(11)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run("Volt Debt · ข้อมูลทดสอบ / Test data · หน้า "), 9)
    _page_number(footer)
    return doc


def heading(doc, th, en, level=1):
    p = doc.add_heading(level=level)
    _font(p.add_run(th))
    sub = doc.add_paragraph()
    _font(sub.add_run(en), 10, color="666666")


def para(doc, th, en=None):
    p = doc.add_paragraph()
    _font(p.add_run(th))
    if en:
        q = doc.add_paragraph()
        _font(q.add_run(en), 10, color="666666")


def steps(doc, items):
    for th, en in items:
        p = doc.add_paragraph(style="List Number")
        _font(p.add_run(th))
        _font(p.add_run("\n" + en), 10, color="666666")


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        _font(cell.paragraphs[0].add_run(text), 10, bold=True)
        _shade(cell, AMBER)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(str(text)), 10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


# --- content ---------------------------------------------------------------------------

def build(passwords, with_passwords):
    doc = new_document()
    title = doc.add_paragraph()
    _font(title.add_run("คู่มือการเข้าสู่ระบบและบัญชีทดสอบ"), 24, bold=True)
    _font(doc.add_paragraph().add_run("Volt Debt — Login and Test Account Guide"), 14, color="666666")
    _font(doc.add_paragraph().add_run("ฉบับ %s · %s%s" % (
        VERSION, date.today().isoformat(), " · มีรหัสผ่าน — ห้ามส่งต่อ" if with_passwords else "")), 11)
    doc.add_page_break()

    heading(doc, "ประวัติเอกสาร", "Document history")
    table(doc, ["ฉบับ · Version", "วันที่ · Date", "เปลี่ยนแปลง · Change"],
          [[VERSION, date.today().isoformat(), "ฉบับแรก · first issue (design 2026-09-13)"]], [3, 3, 11])

    heading(doc, "1. การเข้าสู่ระบบคืออะไร", "1. What the login is")
    para(doc, "บนระบบทดสอบ (staging) มีสองประตูต่อกัน ประตูแรกคือ Google (IAP) ให้เข้าได้เฉพาะบัญชี volt.co.th "
              "ประตูที่สองคือหน้าเข้าสู่ระบบของแอป ที่ใช้ Google หรือบัญชีทดสอบก็ได้",
         "On staging there are two gates. First Google (IAP) lets in volt.co.th accounts only. "
         "Then the app's own login page, with Google or a test account.")
    para(doc, "บนเครื่องของนักพัฒนาไม่มี IAP จะเห็นเฉพาะหน้าเข้าสู่ระบบของแอป",
         "On a developer's laptop there is no IAP; only the app's login page appears.")

    heading(doc, "2. เข้าสู่ระบบด้วย Google", "2. Signing in with Google")
    steps(doc, [
        ("เปิดลิงก์ของระบบ แล้วเลือกบัญชี volt.co.th ของคุณในหน้าของ Google", "Open the link and pick your volt.co.th account on Google's page."),
        ("ในหน้าเข้าสู่ระบบของแอป กดปุ่ม \"เข้าสู่ระบบด้วย Google\"", "On the app's login page, press \"Sign in with Google\"."),
        ("ถ้าเห็น \"บัญชี Google นี้ยังไม่ได้รับสิทธิ์ใช้ระบบ\" แปลว่าอีเมลยังไม่อยู่ในรายชื่อ ให้ขอสิทธิ์จากเจ้าของระบบ (หัวข้อ 5)",
         "\"This Google account is not on the allow-list\" means your email is not listed yet — ask the system owner (section 5)."),
    ])

    heading(doc, "3. เข้าสู่ระบบด้วยบัญชีทดสอบ", "3. Signing in with a test account")
    para(doc, "บัญชีทดสอบทุกบัญชีจะมีป้าย \"บัญชีทดสอบ\" ที่แถบด้านบน และในบันทึกการกระทำ",
         "Every test account shows a \"TEST ACCOUNT\" badge in the top bar and in the action log.")
    table(doc, ["ชื่อผู้ใช้ · Username", "บทบาท · Role", "ทำงานในชื่อ · Acts as", "รหัสผ่าน · Password"],
          [[user, role, officer, passwords.get(user, "—") if with_passwords else MASK]
           for user, role, officer in ACCOUNTS], [3.5, 4.5, 3.5, 5.5])
    para(doc, "tester สลับบทบาทได้จากเมนูชื่อที่มุมขวาบน บัญชีอื่นล็อกไว้บทบาทเดียว",
         "tester can switch roles from the identity menu, top right. The others are locked to one role.")
    para(doc, "ใส่รหัสผ่านผิด 5 ครั้งใน 15 นาที ชื่อผู้ใช้นั้นจะถูกล็อก 15 นาที ถ้าไม่ได้ใช้งาน 30 นาที หรือครบ 8 ชั่วโมง ต้องเข้าสู่ระบบใหม่",
         "5 wrong passwords in 15 minutes locks that username for 15 minutes. After 30 idle minutes, or 8 hours in all, sign in again.")
    para(doc, "การออกจากระบบจะลบคุกกี้ในเบราว์เซอร์นี้เท่านั้น ถ้าคุกกี้ถูกคัดลอกไป จะยังใช้ได้ไม่เกิน 8 ชั่วโมง อย่าใช้บัญชีทดสอบบนเครื่องที่ใช้ร่วมกัน",
         "Signing out clears the cookie in this browser only; a copied cookie stays valid for up to 8 hours. Don't use a test account on a shared computer.")

    heading(doc, "4. รายการตรวจสอบตามบทบาท", "4. Checklist by role")
    para(doc, "สำหรับแต่ละบทบาท ให้ลองทุกแถว \"ควรทำได้\" คือต้องเห็นปุ่มและทำสำเร็จ \"ควรถูกปฏิเสธ\" คือต้องไม่เห็นปุ่ม หรือระบบแจ้งปฏิเสธ",
         "For each role, try every row. \"Should work\" = the control is there and succeeds. \"Should be refused\" = no control, or a refusal message.")
    for role in role_matrix.ROLES:
        heading(doc, "4.%d %s (%s)" % (role_matrix.ROLES.index(role) + 1, ROLE_TH[role], role),
                "Sign in as %s" % next(u for u, r, _o in ACCOUNTS if r == role), level=2)
        rows = []
        for screen_th, screen_en, what_th, what_en, allowed, file, line, _needle in role_matrix.ROWS:
            expect = "ควรทำได้ · should work" if role in allowed else "ควรถูกปฏิเสธ · should be refused"
            rows.append(["%s · %s" % (screen_th, screen_en), "%s · %s" % (what_th, what_en), expect, "☐"])
        table(doc, ["หน้าจอ · Screen", "สิ่งที่ลอง · Try", "ผลที่ควรได้ · Expected", "ผ่าน · Pass"], rows, [3.5, 7, 4.5, 2])

    heading(doc, "5. เพิ่มคนเข้ารายชื่อที่อนุญาต", "5. Adding a person to the allow-list")
    steps(doc, [
        ("แก้ค่า AUTH_ALLOWLIST (JSON) เพิ่มอีเมล บทบาท และชื่อเจ้าหน้าที่ เช่น {\"thip@volt.co.th\": {\"role\": \"OFFICER\", \"officer\": \"thip\"}}",
         "Edit AUTH_ALLOWLIST (JSON): add the email, role and officer name."),
        ("ชื่อเจ้าหน้าที่ต้องมีอยู่ใน /v1/officers ไม่อย่างนั้นระบบจะไม่ยอมเริ่มทำงาน", "The officer must exist in /v1/officers, or the app refuses to start."),
        ("รัน ./scripts/push_secrets.sh --apply แล้ว ./scripts/deploy_gcp.sh", "Run ./scripts/push_secrets.sh --apply, then ./scripts/deploy_gcp.sh."),
    ])

    heading(doc, "6. ปัญหาที่พบบ่อย", "6. Common problems")
    table(doc, ["อาการ · Symptom", "สาเหตุ · Cause", "ทำอย่างไร · Fix"], [
        ["กลับมาหน้าเข้าสู่ระบบเรื่อยๆ · keeps returning to login", "ไม่ได้ใช้งาน 30 นาที หรือครบ 8 ชั่วโมง · idle 30 min or 8 h", "เข้าสู่ระบบใหม่ · sign in again"],
        ["ชื่อผู้ใช้ถูกล็อก · username locked", "รหัสผ่านผิด 5 ครั้ง · 5 wrong passwords", "รอ 15 นาที · wait 15 minutes"],
        ["ไม่มีปุ่ม Google · no Google button", "ยังไม่ได้ตั้งค่า GOOGLE_CLIENT_ID · not configured", "ใช้บัญชีทดสอบ · use a test account"],
        ["หน้าปฏิเสธของ Google ก่อนเข้าแอป · Google refusal before the app", "บัญชีไม่ใช่ volt.co.th (IAP) · not a volt.co.th account", "ใช้บัญชีบริษัท · use your company account"],
        ["ปุ่มที่ควรมีไม่ขึ้น · an expected control is missing", "เข้าสู่ระบบผิดบทบาท · wrong role", "ดูบทบาทที่เมนูชื่อ มุมขวาบน · check the identity menu"],
    ], [5.5, 5.5, 6])
    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--accounts", default=str(Path.home() / ".volt-test-accounts.json"))
    parser.add_argument("--repo-out", default=str(ROOT / "docs" / "guides" / "login-tester-guide.docx"))
    parser.add_argument("--vault-out", default=str(VAULT / "06 - Login Tester Guide (with passwords).docx"))
    args = parser.parse_args()

    accounts_path = Path(args.accounts)
    passwords = json.loads(accounts_path.read_text()) if accounts_path.exists() else {}

    repo_out = Path(args.repo_out)
    repo_out.parent.mkdir(parents=True, exist_ok=True)
    build(passwords, with_passwords=False).save(repo_out)
    print("repo copy (no passwords): %s" % repo_out)

    if passwords:
        vault_out = Path(args.vault_out)
        vault_out.parent.mkdir(parents=True, exist_ok=True)
        build(passwords, with_passwords=True).save(vault_out)
        print("vault copy (with passwords): %s" % vault_out)
    else:
        print("no %s — vault copy not built" % accounts_path)


if __name__ == "__main__":
    main()
