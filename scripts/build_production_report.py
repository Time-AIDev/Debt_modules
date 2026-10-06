#!/usr/bin/env python3
"""Generate the production DOCX from saved evidence; missing evidence stays pending.

Run again after validation/deployment to finalize the same report. Accepts classic
k6 handleSummary objects and the versioned results.metrics machine-readable format.
No network requests, application writes, deployment or Git mutations occur here.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from xml.etree import ElementTree as ET
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PENDING = "รอหลักฐาน / pending"
NAVY, TEAL, MUTED = "16324F", "157F7A", "516579"
SENSITIVE = re.compile(r"(?:token|password|secret|credential|authorization|customer_(?:rows|records|name|email|phone|uid|id)|raw_scope|payload|^email$)", re.I)
SOURCES = [
    ("Grafana k6: metrics", "https://grafana.com/docs/k6/latest/using-k6/metrics/"),
    ("Grafana k6: thresholds", "https://grafana.com/docs/k6/latest/using-k6/thresholds/"),
    ("Grafana k6: end-of-test summaries", "https://grafana.com/docs/k6/latest/results-output/end-of-test/"),
    ("Grafana k6: custom summaries", "https://grafana.com/docs/k6/latest/results-output/end-of-test/custom-summary/"),
    ("Official machine-readable schema", "https://github.com/grafana/k6-summary"),
]


def clean(value):
    """Only presentation-safe evidence; never serialize credentials or raw records."""
    if value is None:
        return PENDING
    if isinstance(value, dict):
        return "; ".join(f"{clean(k)}: {clean(v)}" for k, v in value.items() if not SENSITIVE.search(str(k))) or PENDING
    if isinstance(value, (list, tuple)):
        return "; ".join(clean(v) for v in value) or PENDING
    text = str(value)
    text = re.sub(r"([A-Za-z][A-Za-z0-9+.-]*://)[^/\s:@]+(?::[^/\s@]+)?@", r"\1[redacted]@", text)
    text = re.sub(r"\bBearer\s+[A-Za-z0-9_.-]+", "Bearer [redacted]", text, flags=re.I)
    text = re.sub(r"\beyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+", "[redacted-token]", text)
    text = re.sub(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", "[redacted-email]", text, flags=re.I)
    text = re.sub(r"(?i)((?:password|secret|token|credential)\s*[=:]\s*)[^\s;,]+", r"\1[redacted]", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text[:6000]


def get(obj, *paths, default=None):
    for path in paths:
        current = obj
        for part in path.split("."):
            current = current.get(part) if isinstance(current, dict) else None
        if current is not None:
            return current
    return default


def number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def display(value, suffix="", digits=2):
    value = number(value)
    if value is None:
        return PENDING
    return f"{value:,.{digits}f}{suffix}"


def percent(value):
    value = number(value)
    return display(value * 100, "%") if value is not None else PENDING


def status(value):
    if value is True:
        return "ผ่าน / passed"
    if value is False:
        return "ไม่ผ่าน / failed"
    if value is None or value == "":
        return PENDING
    return clean(value)


def load_json(path, issues):
    if not path.exists():
        issues.append(f"ไม่พบ {path.name}")
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (ValueError, OSError):
        issues.append(f"อ่าน {path.name} ไม่ได้")
        return {}


def junit(path, issues):
    if not path.exists():
        issues.append(f"ไม่พบ {path.name}")
        return {"name": path.name, "status": PENDING}
    try:
        root = ET.parse(path).getroot()
        suites = [root] if root.tag == "testsuite" else root.findall("testsuite")
        # Sum leaf suites; avoid double-counting nested suite aggregate nodes.
        leaves = [s for s in root.iter("testsuite") if not s.findall("testsuite")]
        suites = leaves or suites
        counts = {key: sum(int(s.get(key, "0")) for s in suites) for key in ("tests", "failures", "errors", "skipped")}
        elapsed = sum(float(s.get("time", "0")) for s in suites)
        counts["passed"] = max(0, counts["tests"] - counts["failures"] - counts["errors"] - counts["skipped"])
        counts.update(name=path.name, duration_seconds=elapsed,
                      status="passed" if counts["tests"] > 0 and not counts["failures"] and not counts["errors"] else "failed")
        return counts
    except (ET.ParseError, ValueError, OSError):
        issues.append(f"JUnit ผิดรูปแบบ: {path.name}")
        return {"name": path.name, "status": PENDING}


def metric_map(raw):
    candidate = get(raw, "metrics", "results.metrics", "summary.metrics", default={})
    if isinstance(candidate, list):
        return {row["name"]: row for row in candidate if isinstance(row, dict) and row.get("name")}
    return candidate if isinstance(candidate, dict) else {}


def values(metric):
    if not isinstance(metric, dict):
        return {}
    return metric.get("values") or metric.get("data") or metric


def thresholds(raw, metrics):
    result = []
    for name, metric in metrics.items():
        if not isinstance(metric, dict):
            continue
        entries = metric.get("thresholds", {})
        if isinstance(entries, dict):
            for expression, value in entries.items():
                ok = value.get("ok") if isinstance(value, dict) else value if isinstance(value, bool) else None
                result.append({"metric": name, "expression": expression, "ok": ok})
        elif isinstance(entries, list):
            for entry in entries:
                if isinstance(entry, dict):
                    result.append({"metric": name, "expression": get(entry, "expression", "threshold", "name", default=PENDING),
                                   "ok": get(entry, "ok", "passed", "success")})
    # Versioned summaries can publish the threshold results separately.
    entries = get(raw, "results.thresholds", "thresholds", default=[])
    if isinstance(entries, list):
        for entry in entries:
            if isinstance(entry, dict):
                result.append({"metric": get(entry, "metric", "metricName", default=PENDING),
                               "expression": get(entry, "expression", "threshold", "name", default=PENDING),
                               "ok": get(entry, "ok", "passed", "success")})
    return result


def load_k6(path, issues, metadata=None):
    raw = load_json(path, issues)
    metrics = metric_map(raw)
    latency = values(metrics.get("http_req_duration"))
    requests = values(metrics.get("http_reqs"))
    errors = values(metrics.get("http_req_failed"))
    vus = values(metrics.get("vus"))
    duration = number(get(raw, "state.testRunDurationMs"))
    duration = duration / 1000 if duration is not None else number(get(raw, "config.duration", "metadata.duration_seconds", "metadata.durationSeconds", "duration_seconds"))
    meta = get(raw, "metadata", "meta", default={})
    if not isinstance(meta, dict):
        meta = {}
    if isinstance(metadata, dict):
        meta = dict(meta, **metadata)
    if duration is None:
        duration = number(get(meta, "duration_seconds", "durationSeconds"))
    count = number(requests.get("count"))
    rate = number(requests.get("rate"))
    derived = rate is None and count is not None and duration is not None and duration > 0
    if derived:
        rate = count / duration
    v = number(vus.get("max"))
    if v is None:
        v = number(get(meta, "vus", "virtual_users"))
    threshold_rows = thresholds(raw, metrics)
    ok_values = [r["ok"] for r in threshold_rows]
    outcome = "failed" if False in ok_values else "passed" if ok_values and all(v is True for v in ok_values) else PENDING
    return {"file": path.name, "name": get(meta, "name", "scenario", default=path.stem.replace("-summary", "")),
            "metrics": metrics, "p50": get(latency, "med", "p(50)", "p50"),
            "p95": get(latency, "p(95)", "p95"), "p99": get(latency, "p(99)", "p99"),
            "count": count, "throughput": rate, "derived_throughput": derived,
            "error_rate": errors.get("rate"), "vus": v, "duration": duration,
            "thresholds": threshold_rows, "outcome": outcome, "metadata": meta,
            "dataset": get(meta, "dataset", "dataset_description", default=PENDING),
            "read_write": get(meta, "read_write_mix", "read_write", "workload", default=PENDING)}


def hyperlink(paragraph, title, url):
    relationship = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = clean(title)
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def style_document(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    for name in ("Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Leelawadee UI"
        style.element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Leelawadee UI")
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), "th-TH")
        style.element.get_or_add_rPr().append(language)
        style.font.size = Pt(10 if name in ("Normal", "Body Text") else 24 if name == "Title" else 15 if name == "Heading 1" else 12)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.18
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string(TEAL)
    header = section.header.paragraphs[0]
    header.text = "VOLT COLLECTION  •  PRODUCTION READINESS"
    header.style = "Caption"
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Internal • Evidence-based report  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def table(doc, headers, rows, widths=None):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.style = "Table Grid"
    for cell, title in zip(tab.rows[0].cells, headers):
        cell.text = clean(title)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, row in enumerate(rows):
        cells = tab.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = clean(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            if index % 2 == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEF5F7")
                cell._tc.get_or_add_tcPr().append(shade)
    if widths:
        for row in tab.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()
    return tab


def paragraph(doc, text):
    return doc.add_paragraph(clean(text))


def chart(doc, runs):
    eligible = [run for run in runs if number(run["p95"]) is not None and number(run["throughput"]) is not None]
    if not eligible:
        return
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        names = [str(run["name"]).replace("k6-", "")[:28] for run in eligible]
        fig, axes = plt.subplots(1, 2, figsize=(9, 3.2))
        axes[0].barh(names, [float(r["p95"]) for r in eligible], color="#157F7A")
        axes[0].set_xlabel("HTTP request p95 (ms)")
        axes[1].barh(names, [float(r["throughput"]) for r in eligible], color="#16324F")
        axes[1].set_xlabel("HTTP requests / second")
        for ax in axes:
            ax.grid(axis="x", alpha=0.2)
            ax.set_axisbelow(True)
            ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        data = io.BytesIO()
        fig.savefig(data, format="png", dpi=180, bbox_inches="tight")
        plt.close(fig)
        data.seek(0)
        doc.add_picture(data, width=Inches(6.8))
        paragraph(doc, "กราฟใช้ค่าจาก k6 summary จริงแต่ละรอบ; ความต่างของ VUs, dataset และ workload อาจทำให้เทียบประสิทธิภาพโดยตรงไม่ได้")
    except ImportError:
        paragraph(doc, "ไม่มี matplotlib ใน environment นี้; ตารางตัวเลขยังอ้างอิง evidence โดยตรง")


def sanitized_remote(value):
    if not value:
        return PENDING
    value = str(value)
    if value.startswith("git@"):
        host, _, repo = value[4:].partition(":")
        return f"https://{host}/{repo}"
    if "://" in value:
        parts = urlsplit(value)
        return urlunsplit((parts.scheme, parts.hostname or "", parts.path, "", ""))
    return clean(value)


def validate_docx(path, expected):
    with zipfile.ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError("OOXML archive failed CRC validation")
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        if not required.issubset(archive.namelist()):
            raise ValueError("OOXML package lacks required components")
        for name in archive.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                ET.fromstring(archive.read(name))
        document = ET.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        text = " ".join(node.text or "" for node in document.findall(".//w:t", ns))
        for value in expected:
            if clean(value) not in text:
                raise ValueError(f"Expected report evidence missing: {clean(value)[:80]}")
        if re.search(r"\beyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+|Bearer\s+[A-Za-z0-9_.-]{20,}", text):
            raise ValueError("Credential-like text found in report")
        return {"ooxml_crc": "passed", "xml_parse": "passed", "expected_evidence": "passed",
                "tables": len(document.findall(".//w:tbl", ns)), "text_characters": len(text),
                "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "layout_render": "not performed; OOXML content validation only"}


def build(evidence_dir, output):
    issues = []
    release = load_json(evidence_dir / "release-evidence.json", issues)
    validation = load_json(evidence_dir / "validation.json", issues)
    plan_raw = load_json(evidence_dir / "remediation-plan.json", issues)
    plan = plan_raw if isinstance(plan_raw, list) else get(plan_raw, "findings", "items", default=[])
    plan = plan if isinstance(plan, list) else []
    test_results = [junit(evidence_dir / name, issues) for name in ("backend-tests.xml", "frontend-security-tests.xml")]
    additional = get(validation, "checks", "tests", default=[])
    if isinstance(additional, dict):
        additional = [dict(value, name=name) if isinstance(value, dict) else {"name": name, "status": value} for name, value in additional.items()]
    additional = additional if isinstance(additional, list) else []
    run_meta = get(validation, "k6_runs", "load_runs", default={})
    run_meta = run_meta if isinstance(run_meta, dict) else {}
    runs = [load_k6(path, issues, run_meta.get(path.name) or run_meta.get(path.stem)) for path in sorted(evidence_dir.glob("k6-*-summary.json"))]
    if not runs:
        issues.append("ยังไม่มี k6 summary")
    timestamp = datetime.now(ZoneInfo("Asia/Bangkok")).strftime("%d %B %Y %H:%M ICT")
    doc = Document()
    style_document(doc)
    doc.core_properties.title = "VOLT Collection: Production readiness and k6 evidence"
    doc.core_properties.subject = "Remediation, validation, load testing and deployment evidence"
    doc.core_properties.author = "VOLT Engineering"
    doc.add_paragraph("Production readiness\nand k6 load testing", "Title")
    doc.add_paragraph("รายงานการปรับปรุงและความพร้อมระบบ Collection", "Subtitle")
    paragraph(doc, f"Generated: {timestamp} • เป้าหมาย: collection.volt.co.th / 143.198.207.143")
    paragraph(doc, "รายงานนี้สรุปเฉพาะผลที่มีไฟล์หลักฐานในชุดตรวจ ผลที่ขาดหรือยังทำไม่ได้แสดงเป็น pending และสามารถสร้างเอกสารใหม่หลังงานเสร็จได้")
    deployment = get(release, "deployment", default={})
    git = get(release, "git", default={})
    table(doc, ["Milestone", "Evidence status"], [
        ("Production deployment", status(get(deployment, "status", default=get(release, "deployment_status")))),
        ("Go-live enabled", status(get(deployment, "go_live", "go_live_enabled", default=get(release, "go_live")))),
        ("Git remote push", status(get(git, "pushed", "push_status", "status"))),
        ("k6 summaries present", str(len(runs)) if runs else PENDING),
        ("Remediation register", f"{len(plan)} findings" if plan else PENDING),
    ])
    if issues:
        paragraph(doc, "หลักฐานที่ยังขาด: " + "; ".join(issues))

    doc.add_heading("1. แผนและผลการปรับปรุง / Remediation", level=1)
    paragraph(doc, "ลำดับงาน: แก้ transaction และยอดเงิน → ปิดช่องสิทธิ์และ input → ปรับงานเบื้องหลัง/ไฟล์ → ใช้ runtime policy และ release scripts ร่วมกัน → ตรวจ regression และ k6 → backup/restore และตรวจ domain → บันทึก deployment และ Git push จริง")
    if plan:
        counts = Counter(str(get(row, "status", default="pending")) for row in plan if isinstance(row, dict))
        table(doc, ["Status recorded", "Findings"], sorted(counts.items()))
        for item in plan:
            if not isinstance(item, dict):
                continue
            ident = get(item, "id", "number", default="—")
            title = get(item, "title", "finding", "name", default=PENDING)
            doc.add_heading(f"{ident}. {clean(title)}", level=2)
            table(doc, ["Field", "Evidence"], [
                ("Severity / สถานะ", f"{clean(get(item, 'severity', 'priority'))} • {status(get(item, 'status'))}"),
                ("Changes / การปรับปรุง", clean(get(item, "changes", "remediation", "fix", "action"))),
                ("Verification / การตรวจ", clean(get(item, "verification", "tests", "evidence"))),
                ("Remaining / งานค้าง", "ไม่มีงานค้างระบุไว้ / none recorded" if item.get("remaining") == "" else clean(get(item, "remaining", "limitations", "residual_risk"))),
                ("Code / files", "ไม่ได้ระบุไฟล์ / not enumerated" if item.get("files") == [] else clean(get(item, "files", "sources"))),
            ], widths=[3.0, 14.4])
    else:
        paragraph(doc, "ยังไม่มี remediation-plan.json จึงไม่ยืนยันว่าปิดข้อค้นพบใดแล้ว")

    doc.add_heading("2. ผลตรวจโค้ด / Automated validation", level=1)
    table(doc, ["Evidence", "Status", "Passed", "Failed / errors", "Skipped", "Seconds"], [
        (row["name"], status(row.get("status")), row.get("passed", PENDING),
         row.get("failures", 0) + row.get("errors", 0) if "tests" in row else PENDING,
         row.get("skipped", PENDING), display(row.get("duration_seconds"))) for row in test_results
    ])
    if additional:
        table(doc, ["Check", "Status", "Counts / duration", "Evidence / scope"], [
            (get(row, "name", "check"), status(get(row, "status", "passed")),
             (f"passed={row.get('passed', PENDING)}, failed={row.get('failed', PENDING)}, skipped={row.get('skipped', PENDING)}; "
              f"{display(row.get('duration_seconds'))} s") if any(key in row for key in ("failed", "duration_seconds", "skipped")) else "—",
             get(row, "notes", "detail", "evidence", "result")) for row in additional if isinstance(row, dict)
        ])
    paragraph(doc, "จำนวนที่ผ่านไม่รวม skipped; การผ่าน unit tests ไม่ยืนยัน capacity, ข้อมูลต้นทาง, browser OAuth หรือการกู้คืนหลังเครื่องสูญหาย")

    doc.add_heading("3. ผล k6 / Measured load test", level=1)
    paragraph(doc, "Latency p50/p95/p99 ใช้สถิติที่ k6 export จริง ไม่มีการอนุมาน percentile ที่ขาด Throughput แสดง requests/second; เมื่อคำนวณจาก count ÷ duration จะระบุไว้ ค่า http_req_failed เป็นอัตราตาม expected-response ของสคริปต์ และควรอ่านคู่กับ custom checks โดยเฉพาะ HTTP 409 ที่ตั้งใจทดสอบ")
    if runs:
        table(doc, ["Run", "VUs", "Time (s)", "p50 / p95 / p99 (ms)", "Req/s", "HTTP fail", "Threshold"], [
            (r["name"], display(r["vus"], digits=0), display(r["duration"]),
             f"{display(r['p50'])} / {display(r['p95'])} / {display(r['p99'])}",
             display(r["throughput"]), percent(r["error_rate"]), status(r["outcome"])) for r in runs
        ])
        chart(doc, runs)
        for run in runs:
            doc.add_heading(clean(run["name"]), level=2)
            meta = run["metadata"]
            table(doc, ["Workload field", "Evidence"], [
                ("Summary file", run["file"]), ("Dataset", clean(run["dataset"])),
                ("Read / write mix", clean(run["read_write"])),
                ("Planned VUs / duration", clean(get(meta, "vus", "virtual_users")) + " / " + clean(get(meta, "duration_seconds", "durationSeconds"))),
                ("Target / generator", clean(get(meta, "target", "base_url")) + " / " + clean(get(meta, "generator", "execution"))),
                ("k6 version / start", clean(get(meta, "k6Version", "k6_version")) + " / " + clean(get(meta, "generatedAt", "started_at"))),
                ("HTTP requests", display(run["count"], digits=0)),
                ("Throughput source", "count ÷ measured duration" if run["derived_throughput"] else "k6 exported rate" if run["throughput"] is not None else PENDING),
                ("Exit code / notes", clean(get(meta, "exit_code")) + " / " + clean(get(meta, "notes", "limitations"))),
            ])
            if run["thresholds"]:
                table(doc, ["Metric", "Threshold", "Actual outcome"], [
                    (row["metric"], row["expression"], status(row["ok"])) for row in run["thresholds"]
                ])
            else:
                paragraph(doc, "ไม่มี threshold outcome ใน summary; ไม่ตีความว่าผ่าน acceptance gate")
            custom = [(name, clean(values(metric))) for name, metric in run["metrics"].items()
                      if any(word in name.lower() for word in ("conflict", "duplicate", "idempot", "retry", "error", "write", "read", "lock", "checks", "success", "refusal"))]
            if custom:
                table(doc, ["Business / concurrency metric", "Exported values"], custom)
    else:
        paragraph(doc, "รอผล k6 จริง; ตาราง capacity และ latency ยังไม่พร้อม")
    baselines = [r for r in runs if "baseline" in str(r["name"]).lower() or get(r["metadata"], "phase") == "before"]
    if not baselines:
        paragraph(doc, "Before/after: ไม่มี baseline ก่อนแก้ไขในชุดหลักฐาน จึงไม่อ้างเปอร์เซ็นต์ที่เร็วขึ้นหรือการลด error ก่อนและหลัง")
    else:
        paragraph(doc, "มี baseline อยู่ในตารางผลแต่ละรอบ การเทียบ before/after ใช้ได้เมื่อ dataset, VUs, workload และ target resources เท่ากันเท่านั้น")

    doc.add_heading("4. Dev / staging / production release", level=1)
    defaults_file = evidence_dir / "runtime-defaults.json"
    if defaults_file.exists():
        defaults = load_json(defaults_file, issues)
        table(doc, ["Runtime policy", "Recorded value"], [(k, clean(v)) for k, v in defaults.items() if not SENSITIVE.search(k)])
    table(doc, ["Release field", "Recorded evidence"], [
        ("Domain / IP", clean(get(release, "domain", "deployment.domain")) + " / " + clean(get(release, "ip", "deployment.ip"))),
        ("Timestamp", clean(get(release, "timestamp", "deployment.timestamp"))),
        ("Commit / branch", clean(get(git, "commit", "sha")) + " / " + clean(get(git, "branch"))),
        ("Remote", sanitized_remote(get(git, "remote", "remote_url"))),
        ("Push result", status(get(git, "pushed", "push_status", "status"))),
        ("API digest", clean(get(deployment, "api_image_digest", "api_digest", "images.api"))),
        ("Web digest", clean(get(deployment, "web_image_digest", "web_digest", "images.web"))),
        ("Deployment / go-live", status(get(deployment, "status")) + " / " + status(get(deployment, "go_live", "go_live_enabled"))),
    ])
    checks = get(deployment, "checks", default=[])
    if isinstance(checks, dict):
        checks = [{"name": key, "status": value} for key, value in checks.items()]
    if checks:
        table(doc, ["Deployment check", "Status", "Evidence"], [
            (get(row, "name", "check"), status(get(row, "status", "passed")), get(row, "evidence", "notes", "detail")) for row in checks if isinstance(row, dict)
        ])

    doc.add_heading("5. ข้อจำกัดและงานที่ยังต้องติดตาม / Limits", level=1)
    limitations = [
        "ผลทดสอบมีขอบเขตตาม VUs, ระยะเวลา, dataset และ read/write mix ของแต่ละรอบ ไม่ใช่คำรับรอง workload ทุกชนิด",
        "ถ้า generator อยู่บน VPS เดียวกับระบบ ผลรวมทรัพยากรอาจเปลี่ยน latency/throughput และไม่วัด internet round-trip ของผู้ใช้จริง",
        "การใช้ bearer credentials สำหรับ load test ไม่ทดแทน Google sign-in ผ่าน browser บน production origin จริง",
        "Global mutable-book serialization และ global state version ยังมี capacity ceiling และอาจปฏิเสธการแก้ไขที่ไม่เกี่ยวกันแบบ conservative",
        "Single VPS, off-server backup, alert delivery และ business rules ที่ยังไม่ได้ตัดสิน ต้องประเมินจากหลักฐานและผู้รับผิดชอบจริง",
        "DOCX ผ่านการตรวจ ZIP/XML/table content; ยังไม่มี rendered-page inspection หากไม่มีโปรแกรม Word/LibreOffice ในเครื่องสร้างรายงาน",
    ]
    for source in (validation, release, deployment):
        for key in ("limitations", "risks", "pending", "blocked"):
            extra = get(source, key, default=[])
            limitations.extend(extra if isinstance(extra, list) else [extra] if extra else [])
    for item in dict.fromkeys(clean(item) for item in limitations):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("6. แหล่งหลักฐานและการสร้างซ้ำ / Reproducibility", level=1)
    paragraph(doc, "Command: python scripts/build_production_report.py --evidence-dir reports/production-20261006")
    manifest_paths = sorted(p for p in evidence_dir.iterdir() if p.is_file() and p.suffix in (".json", ".xml") and p.name != "report-validation.json")
    table(doc, ["Evidence file", "SHA-256"], [(p.name, hashlib.sha256(p.read_bytes()).hexdigest()) for p in manifest_paths])
    paragraph(doc, "ข้อมูล metric definitions และการอ่าน threshold อ้างอิงเอกสารผู้พัฒนา k6:")
    for title, url in SOURCES:
        p = doc.add_paragraph()
        hyperlink(p, title, url)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    expected = [row["file"] for row in runs] + [str(len(plan)) + " findings"] if plan else [row["file"] for row in runs]
    for run in runs:
        if number(run["p95"]) is not None:
            expected.append(display(run["p95"]))
        if number(run["throughput"]) is not None:
            expected.append(display(run["throughput"]))
        if number(run["error_rate"]) is not None:
            expected.append(percent(run["error_rate"]))
    checked = validate_docx(output, expected)
    checked.update(generated_at=timestamp, evidence_gaps=issues, k6_runs=len(runs), findings=len(plan))
    (output.parent / "report-validation.json").write_text(json.dumps(checked, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return checked


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence-dir", type=Path, default=ROOT / "reports/production-20261006")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.evidence_dir / "production-readiness-and-k6.docx"
    result = build(args.evidence_dir.resolve(), output.resolve())
    print(json.dumps({"output": str(output), **result}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
